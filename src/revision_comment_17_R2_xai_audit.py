from pathlib import Path
import json
import pandas as pd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'results'/'revision_comments_8_R1_17_18_R2'
S22=ROOT/'supplementary'/'final_revision'/'tables'/'Table_S22_XAI.xlsx'

checks=[]
def check(name, ok, detail=''):
    checks.append((name,bool(ok),detail))
    if not ok:
        raise AssertionError(f'{name}: {detail}')

rep=pd.read_csv(OUT/'OOF_Permutation_Repeated_Summary_Comments_18-R2.csv')
check('Repeated-validation label', set(rep['Validation'])=={'Repeated5x5'}, str(sorted(rep['Validation'].unique())))
check('Repeated held-out fold count', set(rep['N_folds'])=={25}, str(sorted(rep['N_folds'].unique())))

rank=pd.read_csv(OUT/'Rank_Agreement_Excluding_Pr_Comment_8-R1.csv')
check('Pr excluded from rank agreement', 'Pr' not in set(rank['Target']), str(rank['Target'].tolist()))
check('Rank agreement targets', set(rank['Target'])=={'SR','Qm','AF'}, str(rank['Target'].tolist()))

primary=pd.read_csv(OUT/'Table3_Separate_XAI_Comments_17-R2_8-R1.csv')
for forbidden in ['Consensus_score','Composite_score','XAI_consensus_score']:
    check(f'No {forbidden} in Table3', forbidden not in primary.columns, str(primary.columns.tolist()))
check('Separate PI column present', 'OOF_Permutation_Delta_RMSE_mean' in primary.columns)
check('Separate SHAP column present', 'Mean_abs_SHAP' in primary.columns)

wb=load_workbook(S22,data_only=True,read_only=True)
check('S22 separate-rank sheet', 'Separate_Ranks' in wb.sheetnames, str(wb.sheetnames))
check('S22 no consensus analysis sheet', not any('consensus' in s.lower() for s in wb.sheetnames), str(wb.sheetnames))
meta=wb['Metadata']
rows=list(meta.iter_rows(values_only=True))
headers=list(rows[0])
vals=list(rows[1])
md=dict(zip(headers,vals))
check('S22 consensus marked removed', md.get('Consensus_score')=='Removed', str(md))
check('S22 repeated-CV metadata', md.get('Repeated_CV')=='5 folds x 5 repeats', str(md.get('Repeated_CV')))
check('S22 SHAP seed recorded', int(md.get('SHAP_permutation_seed'))==42, str(md.get('SHAP_permutation_seed')))

active_legacy=[p for p in (ROOT/'supplementary'/'tables').glob('*') if 'xai' in p.name.lower() or 'shap' in p.name.lower() or 'permutation' in p.name.lower()]
check('No superseded XAI files in active supplementary/tables', len(active_legacy)==0, str([p.name for p in active_legacy]))

report=OUT/'Comment17_R2_XAI_Audit.txt'
report.write_text('\n'.join([f'{"PASS" if ok else "FAIL"}: {name}' + (f' -- {detail}' if detail else '') for name,ok,detail in checks])+'\n',encoding='utf-8')
print(report.read_text())
