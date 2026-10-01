from pathlib import Path
import hashlib
import json
import pandas as pd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'results'/'revision_comments_8_R1_17_18_R2'
SUP=ROOT/'supplementary'/'final_revision'
SRC=ROOT/'src'/'revision_comments_8_R1_17_18_R2_xai.py'

checks=[]
def check(name, ok, detail=''):
    checks.append((name,bool(ok),detail))
    if not ok:
        raise AssertionError(f'{name}: {detail}')

def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()

code=SRC.read_text(encoding='utf-8')
check('Permutation uses held-out Xte/yte', "permutation_importance(\n            m, Xte, yte" in code)
check('Permutation scoring is held-out RMSE', "scoring='neg_root_mean_squared_error'" in code)
check('Repeated XAI uses 5 repeats', 'N_XAI_REPEATS = 5' in code)
check('Ten permutations per held-out fold', 'N_PERM = 10' in code)
check('SHAP explainer seed fixed', "shap.Explainer(model.predict,masker,algorithm='permutation',seed=SEED)" in code)

rep_det=pd.read_csv(OUT/'OOF_Permutation_Repeated_Detail_Comments_18-R2.csv')
rep_sum=pd.read_csv(OUT/'OOF_Permutation_Repeated_Summary_Comments_18-R2.csv')
lofo_det=pd.read_csv(OUT/'OOF_Permutation_LOFO_Detail_Comments_18-R2.csv')
lofo_sum=pd.read_csv(OUT/'OOF_Permutation_LOFO_Summary_Comments_18-R2.csv')

check('Repeated label is Repeated5x5', set(rep_sum['Validation'])=={'Repeated5x5'}, str(rep_sum['Validation'].unique()))
check('Repeated has 25 held-out folds per target', set(rep_sum['N_folds'])=={25}, str(sorted(rep_sum['N_folds'].unique())))
counts=(rep_det.groupby(['Target','Fold','Feature']).size())
check('Repeated detail has 10 permutations per target/fold/feature', set(counts)=={10}, str(sorted(set(counts))))
check('LOFO uses held-out formulation labels', lofo_det['Heldout_group'].notna().all(), str(lofo_det['Heldout_group'].isna().sum()))
check('LOFO has 8 folds per target', set(lofo_sum['N_folds'])=={8}, str(sorted(lofo_sum['N_folds'].unique())))

expected={'Pr':'Alg_mg_ml','SR':'P_over_TP','Qm':'P_over_TP','AF':'Angle_deg'}
for target,feat in expected.items():
    top=rep_sum[rep_sum['Target']==target].sort_values('OOF_Permutation_rank').iloc[0]['Feature']
    check(f'{target} top OOF permutation feature', top==feat, str(top))

fig_active=SUP/'figures'/'Figure_S5_OOF_Permutation.png'
fig_generated=OUT/'figures'/'FigureS5_OOF_Permutation_Comments_18-R2.png'
check('Active Figure S5 matches generated OOF figure', fig_active.exists() and fig_generated.exists() and sha(fig_active)==sha(fig_generated))

s22=SUP/'tables'/'Table_S22_XAI.xlsx'
xlsx=OUT/'XAI_Analysis_Comments_8-R1_17-18-R2.xlsx'
check('Active Table S22 matches generated XAI workbook', s22.exists() and xlsx.exists() and sha(s22)==sha(xlsx))
wb=load_workbook(xlsx,data_only=True,read_only=True)
md_rows=list(wb['Metadata'].iter_rows(values_only=True))
md=dict(zip(md_rows[0],md_rows[1]))
check('Workbook states held-out permutation method', 'held-out test fold' in str(md.get('Permutation_method','')).lower(), str(md.get('Permutation_method')))
check('Workbook records 5x5 repeated XAI', md.get('Repeated_CV')=='5 folds x 5 repeats', str(md.get('Repeated_CV')))
check('Workbook records 10 permutations', int(md.get('Permutation_repeats_per_fold'))==10, str(md.get('Permutation_repeats_per_fold')))
check('Workbook records SHAP seed 42', int(md.get('SHAP_permutation_seed'))==42, str(md.get('SHAP_permutation_seed')))

meta=json.loads((OUT/'XAI_Metadata_Comments_8-R1_17-18-R2.json').read_text())
check('JSON metadata says held-out permutation', any('held-out test observations' in n for n in meta.get('notes',[])), str(meta.get('notes')))
check('JSON metadata records SHAP seed', meta.get('shap_permutation_seed')==42, str(meta.get('shap_permutation_seed')))

active_legacy=[p for p in (ROOT/'supplementary'/'tables').glob('*') if 'permutation' in p.name.lower() or 'xai' in p.name.lower() or 'shap' in p.name.lower()]
check('No superseded in-sample XAI files in active supplementary/tables', len(active_legacy)==0, str([p.name for p in active_legacy]))

report=OUT/'Comment18_R2_OOF_Permutation_Audit.txt'
report.write_text('\n'.join([f'{"PASS" if ok else "FAIL"}: {name}' + (f' -- {detail}' if detail else '') for name,ok,detail in checks])+'\n',encoding='utf-8')
print(report.read_text())
